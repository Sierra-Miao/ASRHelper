import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["开始时间", "终止时间", "说话人", "内容"]
TIME_FORMAT = "[hh]:mm:ss.000"


def format_time(ms):
    seconds, milliseconds = divmod(round(ms), 1000)
    minutes, seconds = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    return f"{hours:02}:{minutes:02}:{seconds:02}.{milliseconds:03}"


def normalize_segments(result):
    rows = []
    for item in result:
        for segment in item.get("sentence_info", []):
            text = str(segment.get("text", segment.get("sentence", ""))).strip()
            if not text:
                continue
            start, end = float(segment["start"]), float(segment["end"])
            if not math.isfinite(start) or not math.isfinite(end) or start < 0 or end < start:
                raise ValueError("识别结果包含无效时间戳")
            speaker = segment.get("spk")
            label = f"说话人 {int(speaker) + 1}" if speaker is not None else "未知说话人"
            # Excel cells hold at most 32,767 characters. Preserve all text.
            for offset in range(0, len(text), 32000):
                rows.append({"start_ms": round(start), "end_ms": round(end),
                             "speaker": label, "content": text[offset:offset + 32000]})
    if not rows:
        raise ValueError("没有检测到可转写的语音；请确认音频中有清晰的人声")
    return rows


def write_excel(rows, path, title="语音转写"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="205F52")
        cell.font = Font(name="微软雅黑", color="FFFFFF", bold=True, size=11)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    for index, row in enumerate(rows, 2):
        values = [row["start_ms"] / 86400000, row["end_ms"] / 86400000,
                  row["speaker"], row["content"]]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(index, column)
            if isinstance(value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
                # Transcripts and model output are data, never Excel formulas.
                cell.data_type = "s"
            else:
                cell.value = value
                cell.number_format = TIME_FORMAT
            cell.font = Font(name="微软雅黑", size=11, color="253B36")
            cell.alignment = Alignment(vertical="top", wrap_text=column == 4)
            if index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F6F3")
        sheet.row_dimensions[index].height = min(180, max(32, 18 * math.ceil(len(row["content"]) / 44)))
    for index, width in enumerate([19, 19, 18, 95], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def read_excel_table(path):
    """Use the actual exported workbook as the LLM source of truth."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        values = list(sheet.iter_rows(values_only=True))
        if list(values[0]) != HEADERS:
            raise ValueError("Excel 表头不匹配")
        rows = []
        for index, (start, end, speaker, content) in enumerate(values[1:], 1):
            def to_ms(value):
                if hasattr(value, "total_seconds"):
                    return round(value.total_seconds() * 1000)
                if hasattr(value, "hour"):
                    return round((value.hour * 3600 + value.minute * 60 + value.second) * 1000 + value.microsecond / 1000)
                return round(float(value) * 86400000)
            rows.append({"id": index, "开始时间": format_time(to_ms(start)),
                         "终止时间": format_time(to_ms(end)), "说话人": str(speaker or ""),
                         "内容": str(content or "")})
        return rows
    finally:
        book.close()
