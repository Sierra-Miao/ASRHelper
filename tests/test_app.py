import io
import json
import os
import shutil
import subprocess
import tempfile
import time
import unittest
import wave
import zipfile
from pathlib import Path
from unittest.mock import patch

import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import create_app
from backend.config import AISettings, PROMPT, SettingsStore
from backend.excel import HEADERS, normalize_segments, read_excel_table, write_excel
from backend.llm import CorrectionError, correct_workbook, validate_groups
from backend.service import TranscriptionService


ROWS = [
    {"start_ms": 100, "end_ms": 1100, "speaker": "说话人 1", "content": "今天天气"},
    {"start_ms": 1100, "end_ms": 2200, "speaker": "说话人 1", "content": "真步错。"},
    {"start_ms": 2500, "end_ms": 3200, "speaker": "说话人 2", "content": "是的。"},
]
GROUPS = {"groups": [{"source_ids": [1, 2], "content": "今天天气真不错。"}, {"source_ids": [3], "content": "是的。"}]}
TEST_TEMP_ROOT = Path(__file__).resolve().parents[1] / "test-results" / "tmp"
TEST_TEMP_ROOT.mkdir(parents=True, exist_ok=True)


def audio_bytes():
    output = io.BytesIO()
    with wave.open(output, "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(2)
        audio.setframerate(16000)
        audio.writeframes(b"\0\0" * 16000)
    return output.getvalue()


class FakeModel:
    def generate(self, input, **kwargs):
        return [{"text": "今天天气真步错。是的。", "sentence_info": [
            {"start": row["start_ms"], "end": row["end_ms"], "spk": int(row["speaker"][-1]) - 1, "text": row["content"]}
            for row in ROWS]}]


class ExcelAndAI(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=TEST_TEMP_ROOT)
        assert Path(self.temp.name).resolve().is_relative_to(TEST_TEMP_ROOT.resolve())
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "original.xlsx"
        write_excel(ROWS, self.path)

    def test_excel_headers_times_and_formula_safety(self):
        source = [dict(ROWS[0], content='=HYPERLINK("https://example.invalid","test")')]
        write_excel(source, self.path)
        book = load_workbook(self.path)
        self.assertEqual([cell.value for cell in book.active[1]], HEADERS)
        self.assertEqual(book.active["D2"].data_type, "s")
        self.assertEqual(book.active["A2"].number_format, "[hh]:mm:ss.000")
        self.assertEqual(book.active.freeze_panes, "A2")
        book.close()
        table = read_excel_table(self.path)
        self.assertEqual(table[0]["开始时间"], "00:00:00.100")
        self.assertEqual(table[0]["内容"], source[0]["content"])

    def test_invalid_model_output_cannot_drop_or_mix_speakers(self):
        table = read_excel_table(self.path)
        for groups in (
            {"groups": [{"source_ids": [1, 2], "content": "漏掉第三行"}]},
            {"groups": [{"source_ids": [1, 2, 3], "content": "错误合并两个人"}]},
            {"groups": [{"source_ids": [3], "content": "乱序"}]},
            {"groups": [{"source_ids": [True], "content": "无效编号"}]},
        ):
            with self.subTest(groups=groups), self.assertRaises(CorrectionError):
                validate_groups(json.dumps(groups), table, ROWS)
        corrected = validate_groups(json.dumps(GROUPS), table, ROWS)
        self.assertEqual(len(corrected), 2)
        self.assertEqual(corrected[0]["end_ms"], 2200)
        self.assertEqual(corrected[1]["speaker"], "说话人 2")

    def test_chat_mode_sends_excel_contents_and_exact_prompt(self):
        calls = []
        def handler(request):
            calls.append(request)
            body = json.loads(request.content)
            self.assertEqual(request.url.path, "/v1/chat/completions")
            self.assertEqual(request.headers["Authorization"], "Bearer test-key")
            self.assertIn(PROMPT, body["messages"][-1]["content"])
            self.assertIn("真步错", body["messages"][-1]["content"])
            return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(GROUPS)}, "finish_reason": "stop"}]})
        settings = AISettings(base_url="http://mock.invalid/v1", model="test-model", api_key="test-key")
        with httpx.Client(transport=httpx.MockTransport(handler)) as client:
            corrected, warnings = correct_workbook(settings, self.path, ROWS, lambda _: None, client)
        self.assertEqual(corrected[0]["content"], "今天天气真不错。")
        self.assertEqual(warnings, [])
        self.assertEqual(len(calls), 1)

    def test_attachment_upload_responses_and_cleanup(self):
        calls = []
        def handler(request):
            calls.append((request.method, request.url.path))
            if request.url.path == "/v1/files":
                self.assertIn(b"PK\x03\x04", request.content)
                self.assertIn(b"user_data", request.content)
                return httpx.Response(200, json={"id": "file-test"})
            if request.url.path == "/v1/responses":
                body = json.loads(request.content)
                self.assertFalse(body["store"])
                self.assertEqual(body["input"][0]["content"][0]["file_id"], "file-test")
                return httpx.Response(200, json={"status": "completed", "output": [{"content": [{"type": "output_text", "text": json.dumps(GROUPS)}]}]})
            return httpx.Response(200, json={"deleted": True})
        settings = AISettings(base_url="http://mock.invalid/v1", model="test", transport="file")
        with httpx.Client(transport=httpx.MockTransport(handler)) as client:
            corrected, warnings = correct_workbook(settings, self.path, ROWS, lambda _: None, client)
        self.assertEqual(len(corrected), 2)
        self.assertEqual(calls, [("POST", "/v1/files"), ("POST", "/v1/responses"), ("DELETE", "/v1/files/file-test")])
        self.assertFalse(warnings)

    def test_provider_errors_do_not_echo_secret_response(self):
        settings = AISettings(base_url="http://mock.invalid", model="test", api_key="secret")
        with httpx.Client(transport=httpx.MockTransport(lambda request: httpx.Response(401, json={"error": "secret"}))) as client:
            with self.assertRaises(CorrectionError) as error:
                correct_workbook(settings, self.path, ROWS, lambda _: None, client)
        self.assertNotIn("secret", str(error.exception))

    def test_long_tables_are_processed_without_losing_rows(self):
        rows = [dict(ROWS[0], start_ms=i * 1000, end_ms=i * 1000 + 900, content=f"第 {i + 1} 行") for i in range(205)]
        write_excel(rows, self.path)
        calls = []
        def handler(request):
            body = json.loads(request.content)
            table = json.loads(body["messages"][-1]["content"].split("本批次 Excel 表格（JSON 行号仅用于对应原文）：\n")[1])
            calls.append(len(table))
            groups = {"groups": [{"source_ids": [row["id"]], "content": row["内容"]} for row in table]}
            return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(groups)}}]})
        with httpx.Client(transport=httpx.MockTransport(handler)) as client:
            corrected, _ = correct_workbook(AISettings(base_url="http://mock.invalid", model="test"), self.path, rows, lambda _: None, client)
        self.assertEqual(calls, [100, 100, 5])
        self.assertEqual(corrected, rows)


class WebApp(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=TEST_TEMP_ROOT)
        assert Path(self.temp.name).resolve().is_relative_to(TEST_TEMP_ROOT.resolve())
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.service = TranscriptionService(self.root)
        self.service.model = FakeModel()
        self.service.model_state = "ready"
        self.store = SettingsStore(self.root / "settings.json")
        self.client = TestClient(create_app(self.service, self.store))
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)
        self.addCleanup(lambda: self.service.executor.shutdown(wait=True))

    def wait(self, job_id):
        for _ in range(150):
            result = self.client.get(f"/api/jobs/{job_id}").json()
            if result["status"] in {"completed", "failed"}:
                return result
            time.sleep(.05)
        self.fail("Background job did not finish")

    def upload(self, use_ai=False):
        response = self.client.post("/api/jobs", files={"file": ("测试录音.wav", audio_bytes(), "audio/wav")}, data={"speakers": "2", "use_ai": str(use_ai).lower()})
        self.assertEqual(response.status_code, 202, response.text)
        return response.json()["id"]

    def test_upload_download_and_frontend(self):
        page = self.client.get("/")
        self.assertEqual(page.status_code, 200)
        self.assertIn("听笺", page.text)
        self.assertIn("frame-ancestors 'none'", page.headers["content-security-policy"])
        job_id = self.upload()
        job = self.wait(job_id)
        self.assertEqual(job["status"], "completed", job)
        self.assertEqual(len(job["rows"]), 3)
        response = self.client.get(job["downloads"]["original"])
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual([c.value for c in book.active[1]], HEADERS)
        book.close()
        self.assertFalse((self.service.directory(job_id) / "upload.wav").exists())
        self.assertFalse((self.service.directory(job_id) / "decoded.wav").exists())
        self.assertEqual(len(self.client.get("/api/jobs").json()), 1)

    def test_ai_failure_preserves_original_and_retry_produces_both(self):
        self.store.update(AISettings(enabled=True, base_url="http://mock.invalid/v1", model="test", api_key="test-secret"))
        with patch("backend.service.correct_workbook", side_effect=CorrectionError("模拟接口不可用")):
            job_id = self.upload(use_ai=True)
            failed_ai = self.wait(job_id)
        self.assertEqual(failed_ai["status"], "completed")
        self.assertIn("original", failed_ai["downloads"])
        self.assertNotIn("corrected", failed_ai["downloads"])
        self.assertIn("不可用", failed_ai["ai_error"])
        corrected = validate_groups(json.dumps(GROUPS), read_excel_table(self.service.directory(job_id) / "original.xlsx"), ROWS)
        with patch("backend.service.correct_workbook", return_value=(corrected, [])):
            self.assertEqual(self.client.post(f"/api/jobs/{job_id}/correct").status_code, 202)
            job = self.wait(job_id)
        self.assertEqual(len(job["corrected_rows"]), 2)
        archive = self.client.get(job["downloads"]["bundle"])
        with zipfile.ZipFile(io.BytesIO(archive.content)) as bundle:
            self.assertEqual(set(bundle.namelist()), {"原始识别.xlsx", "AI校对.xlsx"})
        self.assertNotIn("test-secret", (self.service.directory(job_id) / "job.json").read_text(encoding="utf-8"))

    def test_api_key_persists_across_restart_but_is_never_returned(self):
        response = self.client.put("/api/settings", json={"enabled": True, "base_url": "http://localhost:1234/v1", "model": "local", "api_key": "test-secret"})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["has_api_key"])
        self.assertNotIn("test-secret", response.text)
        self.assertEqual(json.loads(self.store.path.read_text(encoding="utf-8"))["api_key"], "test-secret")
        self.client.put("/api/settings", json={"enabled": True, "base_url": "http://localhost:1234/v1", "model": "local", "api_key": ""})
        self.assertEqual(self.store.snapshot().api_key, "test-secret")
        restored = SettingsStore(self.store.path)
        self.assertTrue(restored.snapshot().enabled)
        self.assertEqual(restored.snapshot().api_key, "test-secret")
        self.assertEqual(restored.snapshot().base_url, "http://localhost:1234/v1")
        self.assertNotIn("test-secret", json.dumps(restored.public()))
        cleared = self.client.put("/api/settings", json={"enabled": False, "keep_key": False})
        self.assertFalse(cleared.json()["has_api_key"])
        self.assertEqual(SettingsStore(self.store.path).snapshot().api_key, "")

    def test_bad_uploads_and_cross_origin_requests(self):
        self.assertEqual(self.client.post("/api/jobs", files={"file": ("bad.exe", b"bad")}).status_code, 400)
        self.assertEqual(self.client.post("/api/jobs", files={"file": ("empty.wav", b"")}).status_code, 400)
        self.assertEqual(self.client.post("/api/jobs", files={"file": ("test.wav", audio_bytes())}, data={"speakers": "0"}).status_code, 400)
        self.assertEqual(self.client.put("/api/settings", json={}, headers={"origin": "https://evil.invalid"}).status_code, 403)
        self.assertEqual(self.client.get("/api/jobs/does-not-exist").status_code, 404)


@unittest.skipUnless(os.name == "nt", "Windows batch launcher")
class LauncherConsent(unittest.TestCase):
    def test_missing_environment_is_not_installed_without_consent(self):
        with tempfile.TemporaryDirectory(dir=TEST_TEMP_ROOT) as directory:
            root = Path(directory).resolve()
            self.assertTrue(root.is_relative_to(TEST_TEMP_ROOT.resolve()))
            shutil.copyfile(Path(__file__).resolve().parents[1] / "launcher.ps1", root / "launcher.ps1")
            command = ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(root / "launcher.ps1")]
            checked = subprocess.run(command + ["-CheckOnly"], capture_output=True, timeout=20)
            self.assertEqual(checked.returncode, 1)
            self.assertFalse(json.loads(checked.stdout)["ready"])
            unattended = subprocess.run(command + ["-NoInstall"], capture_output=True, timeout=20)
            self.assertEqual(unattended.returncode, 1)
            declined = subprocess.run(command, input=b"N\r\n", capture_output=True, timeout=20)
            self.assertEqual(declined.returncode, 2)
            self.assertFalse((root / ".venv").exists())
            self.assertFalse((root / ".tools").exists())


if __name__ == "__main__":
    unittest.main()
