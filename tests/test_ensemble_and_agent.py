import io
import json
import tempfile
import time
import unittest
import wave
import weakref
import zipfile
from pathlib import Path
from unittest.mock import patch

import httpx
from fastapi.testclient import TestClient
from backend.app import create_app
from backend.config import AISettings, SettingsStore, PROMPT
from backend.ensemble import SequentialEnsemble
from backend.excel import read_excel_table, write_excel
from backend.llm import correct_workbook
from backend.paragraphs import local_paragraphs
from backend.review import CorrectionError, validate_groups
from backend.service import TranscriptionService

ROOT = Path(__file__).resolve().parents[1]
TEMP = ROOT / "test-results" / "tmp"
TEMP.mkdir(parents=True, exist_ok=True)
ROWS = [{"start_ms": 100, "end_ms": 400, "speaker": "说话人 1", "content": "我们明天开会。"},
        {"start_ms": 450, "end_ms": 700, "speaker": "说话人 1", "content": "讨论部署方案。"},
        {"start_ms": 750, "end_ms": 900, "speaker": "说话人 2", "content": "我会参加。"}]
RESULT = [{"sentence_info": [{"start": r["start_ms"], "end": r["end_ms"], "spk": int(r["speaker"][-1])-1, "text": r["content"]} for r in ROWS]}]
GROUPS = {"groups": [{"source_ids": [1, 2], "content": "我们明天开会，讨论部署方案。"}, {"source_ids": [3], "content": "我会参加。"}]}


class Fake:
    kwargs = {"device": "cpu"}
    def generate(self, **kwargs):
        return RESULT


def audio():
    buffer = io.BytesIO()
    with wave.open(buffer, "wb") as handle:
        handle.setparams((1, 2, 16000, 0, "NONE", "not compressed"))
        handle.writeframes(b"\0\0" * 16000)
    return buffer.getvalue()


class EnsembleAndParagraphs(unittest.TestCase):
    def test_paragraphs_merge_complete_sentences_without_splitting_or_estimating_times(self):
        merged = local_paragraphs(ROWS)
        self.assertEqual(len(merged), 2)
        self.assertEqual((merged[0]["start_ms"], merged[0]["end_ms"]), (100, 700))
        self.assertEqual(merged[0]["content"], ROWS[0]["content"] + ROWS[1]["content"])
        self.assertEqual(ROWS[0]["end_ms"], 400)
        table = [{"id": i+1} for i in range(3)]
        output = validate_groups(json.dumps(GROUPS), table, ROWS)
        self.assertEqual((output[0]["start_ms"], output[0]["end_ms"]), (100, 700))
        with self.assertRaises(CorrectionError):
            validate_groups(json.dumps({"groups": [{"source_ids": [1], "content": "拆1"}, {"source_ids": [1], "content": "拆2"}, *GROUPS["groups"][1:]]}), table, ROWS)

    def test_dual_pass_releases_primary_before_nano_and_saves_primary_first(self):
        events, weak = [], []
        def loader(profile):
            if profile == "nano":
                self.assertIsNone(weak[0]())
                self.assertIn("saved", events)
            model = Fake()
            if profile == "sensevoice": weak.append(weakref.ref(model))
            events.append(profile)
            return model
        pipeline = SequentialEnsemble(loader=loader)
        result = pipeline.generate("audio.wav", primary_callback=lambda _: events.append("saved"))
        self.assertEqual(events, ["sensevoice", "saved", "nano"])
        self.assertEqual(result, RESULT)
        self.assertEqual(pipeline.reference_result, RESULT)

    def test_secondary_failure_retains_primary_and_reports_degraded_mode(self):
        def loader(profile):
            if profile == "nano": raise MemoryError()
            return Fake()
        pipeline = SequentialEnsemble(loader=loader)
        with self.assertLogs(level="ERROR"):
            self.assertEqual(pipeline.generate("audio.wav"), RESULT)
        self.assertFalse(pipeline.reference_result)
        self.assertIn("内存不足", pipeline.warnings[0])


class AgentHandoff(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=TEMP)
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        with patch("backend.service.hardware_plan", return_value={"cpu_threads": 2}):
            self.service = TranscriptionService(self.root)
        self.service.profile = "ensemble"
        self.service.model = SequentialEnsemble(loader=lambda _: Fake())
        self.addCleanup(lambda: self.service.executor.shutdown(wait=True))
        store = SettingsStore(self.root / "settings.json")
        store.update(AISettings(enabled=True, transport="agent", api_key="secret-key-never-export"))
        self.client = TestClient(create_app(self.service, store))
        self.client.__enter__()
        self.addCleanup(self.client.__exit__, None, None, None)

    def job(self):
        response = self.client.post("/api/jobs", files={"file": ("sample.wav", audio(), "audio/wav")}, data={"use_ai": "true", "speakers": "2"})
        self.assertEqual(response.status_code, 202, response.text)
        job_id = response.json()["id"]
        for _ in range(200):
            job = self.client.get(f"/api/jobs/{job_id}").json()
            if job["status"] not in {"queued", "decoding", "loading", "transcribing", "exporting", "correcting"}: return job
            time.sleep(.02)
        self.fail("job did not complete")

    def test_dual_reference_agent_bundle_direct_excel_handoff(self):
        job = self.job()
        self.assertTrue(job["agent_pending"], job)
        self.assertEqual(job["reference_status"], "completed")
        self.assertEqual(len(job["paragraph_rows"]), 2)
        self.assertEqual(len(job["rows"]), 3)
        package = self.client.get(job["downloads"]["agent"])
        with zipfile.ZipFile(io.BytesIO(package.content)) as archive:
            self.assertIn("Nano参考.xlsx", archive.namelist())
            self.assertIn("主稿.xlsx", archive.namelist())
            self.assertIn("skills/asr-paragraph-review/SKILL.md", archive.namelist())
            self.assertFalse(any(name.endswith(".json") for name in archive.namelist()))
            self.assertFalse(any(b"secret-key-never-export" in archive.read(name) for name in archive.namelist()))
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(archive.read("主稿.xlsx")))
            self.assertEqual(book.active.max_row, 4)
            book.close()
        self.assertEqual(self.client.get(job["downloads"]["paragraphs"]).status_code, 200)
        self.assertEqual(self.client.get("/api/agent-skill").status_code, 200)
        self.assertTrue(self.client.post("/api/settings/test").json()["ok"])

    def test_api_receives_only_time_matched_nano_reference(self):
        path = self.root / "original.xlsx"
        write_excel(ROWS, path)
        refs = [{"start_ms": 0, "end_ms": 1000, "speaker": "IGNORED", "content": "Nano reference in range"},
                {"start_ms": 5000, "end_ms": 6000, "speaker": "IGNORED", "content": "outside clip"}]
        def handler(request):
            body = json.loads(request.content)
            prompt = body["messages"][-1]["content"]
            self.assertIn("Nano reference in range", prompt)
            self.assertNotIn("outside clip", prompt)
            self.assertNotIn("IGNORED", prompt)
            return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(GROUPS)}}]})
        with httpx.Client(transport=httpx.MockTransport(handler)) as client:
            output, _ = correct_workbook(AISettings(base_url="http://mock.invalid/v1", model="mock"), path, ROWS, lambda _: None, client, refs)
        self.assertEqual(output[0]["end_ms"], 700)


if __name__ == "__main__": unittest.main()
